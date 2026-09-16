# modules/exporter.py
"""Modul pembuatan berkas ekspor Excel profesional multi-sheet dengan openpyxl."""

import io
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd


def konversi_ke_excel(df: pd.DataFrame) -> bytes:
    """Mengonversi DataFrame ke bytes Excel dengan styling rapi dan pemisahan per semester."""
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # Sheet gabungan seluruh nilai
        df.to_excel(writer, index=False, sheet_name="Semua Nilai")

        # Sheet per semester
        daftar_semester = sorted(df["semester"].unique())
        for sem in daftar_semester:
            df_sem = df[df["semester"] == sem]
            df_sem.to_excel(writer, index=False, sheet_name=f"Semester {sem}")

        workbook = writer.book

        # Palet Warna & Format
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="Calibri", size=10)
        border_tipis = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        for sheetname in workbook.sheetnames:
            worksheet = workbook[sheetname]

            # Format Header
            for col_num in range(1, worksheet.max_column + 1):
                cell = worksheet.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_center
                cell.border = border_tipis
            worksheet.row_dimensions[1].height = 25

            # Format Baris Data & Zebra Striping
            for row_idx in range(2, worksheet.max_row + 1):
                worksheet.row_dimensions[row_idx].height = 20
                bg_color = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
                row_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

                for col_idx in range(1, worksheet.max_column + 1):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    cell.font = data_font
                    cell.fill = row_fill
                    cell.border = border_tipis

                    header_val = str(worksheet.cell(row=1, column=col_idx).value or "").lower()
                    if "mata" in header_val:
                        cell.alignment = align_left
                    else:
                        cell.alignment = align_center

            # Auto-fit Lebar Kolom
            for col in worksheet.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    return output.getvalue()
